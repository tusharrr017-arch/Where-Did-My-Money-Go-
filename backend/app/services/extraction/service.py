import csv
import io
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def extract_from_csv(contents: bytes) -> str:
    text = contents.decode("utf-8-sig")

    reader = csv.DictReader(io.StringIO(text))

    rows = list(reader)

    return str(rows)


def extract_from_xlsx(contents: bytes) -> str:
    workbook = load_workbook(
        filename=io.BytesIO(contents),
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    rows = []

    for row in sheet.iter_rows(values_only=True):
        rows.append(list(row))

    workbook.close()

    return str(rows)


def extract_from_docx(contents: bytes) -> str:
    document = Document(io.BytesIO(contents))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n".join(paragraphs)


def extract_from_pdf(contents: bytes) -> str:
    reader = PdfReader(io.BytesIO(contents))

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages)


def extract_document(filename: str, contents: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension == ".csv":
        return extract_from_csv(contents)

    if extension == ".xlsx":
        return extract_from_xlsx(contents)

    if extension == ".docx":
        return extract_from_docx(contents)

    if extension == ".pdf":
        return extract_from_pdf(contents)

    raise ValueError(
        "Unsupported file type. "
        "Supported formats: CSV, PDF, XLSX, DOCX."
    )