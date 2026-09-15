from __future__ import annotations

import csv
import io
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.ingestion.dates import infer_statement_year
from app.services.ingestion.types import ExtractedDocument


def extract_structured(filename: str, contents: bytes) -> ExtractedDocument:
    extension = Path(filename).suffix.lower()

    if extension == ".csv":
        return _extract_csv(contents)
    if extension == ".xlsx":
        return _extract_xlsx(contents)
    if extension == ".docx":
        return _extract_docx(contents)
    if extension == ".pdf":
        return _extract_pdf(contents)

    raise ValueError(
        "Unsupported file type. Supported formats: CSV, PDF, XLSX, DOCX."
    )


def _extract_csv(contents: bytes) -> ExtractedDocument:
    text = contents.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]
    statement_year = infer_statement_year(text)
    return ExtractedDocument(
        file_type="csv",
        table_rows=rows,
        text_lines=text.splitlines(),
        statement_year=statement_year,
    )


def _extract_xlsx(contents: bytes) -> ExtractedDocument:
    workbook = load_workbook(
        filename=io.BytesIO(contents),
        read_only=True,
        data_only=True,
    )
    sheet = workbook.active
    rows = [
        [str(cell).strip() if cell is not None else "" for cell in row]
        for row in sheet.iter_rows(values_only=True)
        if any(cell is not None and str(cell).strip() for cell in row)
    ]
    workbook.close()
    joined = "\n".join(" | ".join(row) for row in rows)
    return ExtractedDocument(
        file_type="xlsx",
        table_rows=rows,
        text_lines=joined.splitlines(),
        statement_year=infer_statement_year(joined),
    )


def _extract_docx(contents: bytes) -> ExtractedDocument:
    document = Document(io.BytesIO(contents))
    lines: list[str] = []
    table_rows: list[list[str]] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                table_rows.append(cells)
                lines.append(" | ".join(cells))

    joined = "\n".join(lines)
    return ExtractedDocument(
        file_type="docx",
        text_lines=lines,
        table_rows=table_rows,
        statement_year=infer_statement_year(joined),
    )


def _extract_pdf(contents: bytes) -> ExtractedDocument:
    reader = PdfReader(io.BytesIO(contents))
    lines: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        for line in text.splitlines():
            cleaned = line.strip()
            if cleaned:
                lines.append(cleaned)
    joined = "\n".join(lines)
    return ExtractedDocument(
        file_type="pdf",
        text_lines=lines,
        statement_year=infer_statement_year(joined),
    )
